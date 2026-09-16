"""
Builds downloadable XLSX and PDF reports from an already-computed analysis
payload (the same JSON shape returned by /api/analyze and
/api/analyze/upload). Stamps each export with the report-generation
timestamp and, when available, a live price snapshot distinct from the
last historical close.
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

HEADER_FILL = "1F2A44"
FLAG_FILL = "FBE3DC"
NORMAL_FILL = "FFFFFF"
ACCENT_HEX = colors.HexColor("#B45A3C")
FLAG_HEX = colors.HexColor("#F2795C")
HEADER_HEX = colors.HexColor(f"#{HEADER_FILL}")
FLAG_ROW_HEX = colors.HexColor(f"#{FLAG_FILL}")


def _fmt_price(v: Optional[float]) -> str:
    if v is None:
        return "—"
    if v < 1:
        return f"${v:.4f}"
    if v < 10:
        return f"${v:.3f}"
    return f"${v:,.2f}"


def _fmt_date(iso: Optional[str]) -> str:
    if not iso:
        return "—"
    return datetime.fromisoformat(iso).strftime("%b %d, %Y")


def _fmt_pct(v: Optional[float]) -> str:
    if v is None:
        return "—"
    return f"{v:,.1f}%"


def _meta_lines(payload: dict, live_quote: Optional[dict]) -> list[str]:
    lines = [
        f"Ticker: {payload['ticker']}",
        f"Data source: {'Live Yahoo Finance' if payload['source'] == 'yahoo' else 'Uploaded CSV'}",
        f"Historical data range: {_fmt_date(payload['range_start'])} \u2013 {_fmt_date(payload['range_end'])} "
        f"({payload['trading_days']:,} sessions)",
        f"Last historical close: {_fmt_price(payload['latest_close'])} on {_fmt_date(payload['range_end'])}",
    ]
    if live_quote and live_quote.get("price") is not None:
        lines.append(f"Live price snapshot: {_fmt_price(live_quote['price'])} ({live_quote.get('note', 'live quote')})")
    else:
        note = (live_quote or {}).get("error", "not available in this environment")
        lines.append(f"Live price snapshot: unavailable ({note})")
    lines.append(f"Report generated: {datetime.now().strftime('%b %d, %Y %I:%M %p')}")
    return lines


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

def build_xlsx(payload: dict, live_quote: Optional[dict]) -> bytes:
    wb = Workbook()

    # ---- Summary sheet ----
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"{payload['ticker']} \u2014 SignalLedger"
    ws["A1"].font = Font(size=16, bold=True)
    for i, line in enumerate(_meta_lines(payload, live_quote), start=3):
        ws.cell(row=i, column=1, value=line)
    ws.column_dimensions["A"].width = 70

    # ---- Yearly sheet ----
    ws2 = wb.create_sheet("Yearly")
    headers = ["Year", "High", "High Date", "Low", "Low Date", "Sequence OK",
               "% Diff (Original)", "Revised High (next yr)", "% Diff (Revised)", "Note"]
    ws2.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws2.cell(row=1, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
        c.alignment = Alignment(horizontal="center")

    for row_idx, y in enumerate(payload["yearly"], start=2):
        label = y["label"] + ("" if y["is_complete"] else " (partial)")
        ws2.append([
            label,
            y["high"]["price"], _fmt_date(y["high"]["date"]),
            y["low"]["price"], _fmt_date(y["low"]["date"]),
            "Yes" if y["sequence_ok"] else "No \u2014 low after high",
            y["pct_diff"],
            y["revised_high"]["price"] if y.get("revised_high") else None,
            y["revised_pct_diff"],
            y.get("revised_note") or "",
        ])
        if not y["sequence_ok"]:
            for col in range(1, len(headers) + 1):
                ws2.cell(row=row_idx, column=col).fill = PatternFill("solid", fgColor=FLAG_FILL)

    for col_letter, width in zip("ABCDEFGHIJ", [16, 12, 14, 12, 14, 20, 16, 20, 16, 42]):
        ws2.column_dimensions[col_letter].width = width
    ws2.freeze_panes = "A2"

    # ---- Monthly sheet ----
    ws3 = wb.create_sheet("Monthly")
    m_headers = ["Month", "High", "High Date", "Low", "Low Date", "Sequence OK", "% Diff"]
    ws3.append(m_headers)
    for col in range(1, len(m_headers) + 1):
        c = ws3.cell(row=1, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
        c.alignment = Alignment(horizontal="center")

    for row_idx, m in enumerate(payload["monthly"], start=2):
        label = m["label"] + ("" if m["is_complete"] else " (partial)")
        ws3.append([
            label,
            m["high"]["price"], _fmt_date(m["high"]["date"]),
            m["low"]["price"], _fmt_date(m["low"]["date"]),
            "Yes" if m["sequence_ok"] else "No",
            m["pct_diff"],
        ])
        if not m["sequence_ok"]:
            for col in range(1, len(m_headers) + 1):
                ws3.cell(row=row_idx, column=col).fill = PatternFill("solid", fgColor=FLAG_FILL)

    for col_letter, width in zip("ABCDEFG", [16, 12, 14, 12, 14, 14, 12]):
        ws3.column_dimensions[col_letter].width = width
    ws3.freeze_panes = "A2"

    # ---- Best moves sheet ----
    ws4 = wb.create_sheet("Best Sequential Move")
    ws4.append(["Window", "Low", "Low Date", "High", "High Date", "% Diff"])
    for col in range(1, 7):
        c = ws4.cell(row=1, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
    for label, move in [("Current year", payload.get("best_move_current_year")),
                         ("Full history on file", payload.get("best_move_overall"))]:
        if move:
            ws4.append([label, move["low"]["price"], _fmt_date(move["low"]["date"]),
                        move["high"]["price"], _fmt_date(move["high"]["date"]), move["pct_diff"]])
        else:
            ws4.append([label, "n/a", "", "", "", ""])
    for col_letter, width in zip("ABCDEF", [22, 12, 14, 12, 14, 12]):
        ws4.column_dimensions[col_letter].width = width

    # ---- Daily prices sheet (full history, for reference / your own charts) ----
    history = payload.get("price_history") or []
    if history:
        ws5 = wb.create_sheet("Daily Prices")
        ws5.append(["Date", "Close", "High", "Low"])
        for col in range(1, 5):
            c = ws5.cell(row=1, column=col)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=HEADER_FILL)
        for row in history:
            ws5.append([_fmt_date(row["date"]), row["close"], row["high"], row["low"]])
        for col_letter, width in zip("ABCD", [16, 12, 12, 12]):
            ws5.column_dimensions[col_letter].width = width
        ws5.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def build_pdf(payload: dict, live_quote: Optional[dict]) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=letter,
        topMargin=0.6 * inch, bottomMargin=0.6 * inch,
        leftMargin=0.6 * inch, rightMargin=0.6 * inch,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleX", parent=styles["Title"], fontSize=20, spaceAfter=4)
    meta_style = ParagraphStyle("Meta", parent=styles["Normal"], fontSize=9.5, textColor=colors.HexColor("#444444"), leading=14)
    h2_style = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=13, spaceBefore=16, spaceAfter=6)
    note_style = ParagraphStyle("Note", parent=styles["Normal"], fontSize=8.5, textColor=colors.HexColor("#666666"), spaceAfter=8)

    story = [Paragraph(f"{payload['ticker']} \u2014 SignalLedger", title_style)]
    for line in _meta_lines(payload, live_quote):
        story.append(Paragraph(line, meta_style))

    # Yearly table
    story.append(Paragraph("Yearly range", h2_style))
    story.append(Paragraph(
        "Rows marked FLAGGED had their low occur after their high (a decline year); "
        "the revised % diff uses the following year's actual high instead.", note_style,
    ))
    y_rows = [["Year", "High", "High Date", "Low", "Low Date", "% Diff (orig.)", "% Diff (revised)"]]
    y_flag_rows = []
    for i, y in enumerate(payload["yearly"], start=1):
        label = y["label"] + ("" if y["is_complete"] else "*")
        revised = _fmt_pct(y["revised_pct_diff"]) if not y["sequence_ok"] else "\u2014"
        y_rows.append([
            label, _fmt_price(y["high"]["price"]), _fmt_date(y["high"]["date"]),
            _fmt_price(y["low"]["price"]), _fmt_date(y["low"]["date"]),
            _fmt_pct(y["pct_diff"]), revised,
        ])
        if not y["sequence_ok"]:
            y_flag_rows.append(i)

    y_table = Table(y_rows, repeatRows=1, hAlign="LEFT")
    y_style = [
        ("BACKGROUND", (0, 0), (-1, 0), HEADER_HEX),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("ALIGN", (0, 0), (0, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F7F9")]),
    ]
    for r in y_flag_rows:
        y_style.append(("BACKGROUND", (0, r), (-1, r), FLAG_ROW_HEX))
    y_table.setStyle(TableStyle(y_style))
    story.append(y_table)
    if any(not y["is_complete"] for y in payload["yearly"]):
        story.append(Paragraph("* partial year (not yet complete)", note_style))

    # Best sequential move
    story.append(Paragraph("Best sequential move", h2_style))
    story.append(Paragraph(
        "Largest gain actually capturable buying at a low and selling at a later high, in chronological order.",
        note_style,
    ))
    mv_rows = [["Window", "Low", "Low Date", "High", "High Date", "% Diff"]]
    for label, move in [("Current year", payload.get("best_move_current_year")),
                         ("Full history on file", payload.get("best_move_overall"))]:
        if move:
            mv_rows.append([label, _fmt_price(move["low"]["price"]), _fmt_date(move["low"]["date"]),
                             _fmt_price(move["high"]["price"]), _fmt_date(move["high"]["date"]),
                             _fmt_pct(move["pct_diff"])])
        else:
            mv_rows.append([label, "n/a", "", "", "", ""])
    mv_table = Table(mv_rows, hAlign="LEFT")
    mv_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), HEADER_HEX),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
    ]))
    story.append(mv_table)

    # Monthly table
    story.append(Paragraph("Trailing months", h2_style))
    m_rows = [["Month", "High", "High Date", "Low", "Low Date", "% Diff"]]
    m_flag_rows = []
    for i, m in enumerate(payload["monthly"], start=1):
        label = m["label"] + ("" if m["is_complete"] else "*")
        m_rows.append([
            label, _fmt_price(m["high"]["price"]), _fmt_date(m["high"]["date"]),
            _fmt_price(m["low"]["price"]), _fmt_date(m["low"]["date"]),
            _fmt_pct(m["pct_diff"]),
        ])
        if not m["sequence_ok"]:
            m_flag_rows.append(i)

    m_table = Table(m_rows, repeatRows=1, hAlign="LEFT")
    m_style = [
        ("BACKGROUND", (0, 0), (-1, 0), HEADER_HEX),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F7F9")]),
    ]
    for r in m_flag_rows:
        m_style.append(("BACKGROUND", (0, r), (-1, r), FLAG_ROW_HEX))
    m_table.setStyle(TableStyle(m_style))
    story.append(m_table)

    doc.build(story)
    return buf.getvalue()
